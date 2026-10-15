"""End-to-end checks that run the real Streamlit app headlessly.
Run with:  python -m pytest tests/ -q   (or: python tests/test_app.py)"""
import io
import os
import sys

from streamlit.testing.v1 import AppTest
from docx import Document

APP = os.path.join(os.path.dirname(__file__), "..", "app.py")

BILL = [
    {"description": "Complete Haemogram CBC", "code": "LB012", "claimed_amount": 300},
    {"description": "Haemoglobin (Hb)", "code": "LB007", "claimed_amount": 50},
    {"description": "Total Leucocytic Count", "code": "", "claimed_amount": 50},   # matched from the description
    {"description": "Platelet Count", "code": "LB013", "claimed_amount": 100},
    {"description": "Kidney Function Test", "code": "LB123", "claimed_amount": 650},  # CGHS rate is 500
    {"description": "Serum Creatinine", "code": "LB058", "claimed_amount": 100},
    {"description": "HbA1c", "code": "LB122", "claimed_amount": 300},
    {"description": "HbA1c again", "code": "LB122", "claimed_amount": 300},
]


def _app(items=None):
    at = AppTest.from_file(APP, default_timeout=60)
    if items is not None:
        at.session_state["line_items"] = items
    return at.run()


def _text(at, label):
    return [t for t in at.text_input if t.label == label][0]


def _doc_text(data):
    d = Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


def test_app_runs_clean_and_does_not_pay_packaged_tests_twice():
    at = _app(BILL)
    assert not at.exception
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["Total claimed"] == "₹1,850"
    # CBC 300 + KFT 500 (capped) + HbA1c 300 (the repeat, Hb/TLC/Platelet and creatinine are not paid again)
    assert metrics["Total admissible (as per CGHS rates)"] == "₹1,100"
    sheet = [d.value for d in at.dataframe if "Result" in d.value.columns][0]
    status = dict(zip(sheet["Particulars"].str[:12], sheet["Result"]))
    assert status["Haemoglobin "] == "INCLUDED IN PACKAGE"
    assert list(sheet["Result"]).count("REPEAT - COUNTED ONCE") == 1


def test_gemini_dropdown_lists_current_models():
    at = _app()
    at.sidebar.selectbox[0].set_value("Google (Gemini)").run()
    picker = [s for s in at.sidebar.selectbox if s.key == "gemini_model"][0]
    assert picker.value == "gemini-3.8-flash"
    shown = " | ".join(picker.options)          # AppTest reports the display labels, which start with the model ID
    for m in ("gemini-4-argon", "gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite", "gemini-flash-latest"):
        assert m in shown
    assert "gemini-2.0-flash" not in shown                    # shut down by Google
    picker.set_value("Other (type a model name)...").run()
    assert [t for t in at.sidebar.text_input if t.key == "gemini_model_custom"]


def test_office_details_flow_into_both_documents():
    at = _app(BILL)
    at.text_input(key="office_division").set_value("North Division")
    at.text_input(key="office_address").set_value("Dak Bhawan, Delhi - 110001")
    at.text_input(key="office_memo_no_prefix").set_value("A-1/NORTH/MR")
    at.text_input(key="office_prepared_by").set_value("Sr. Accountant")
    at.text_area(key="office_copy_to").set_value("The Director, Postal Accounts\nO/c")
    at.run()
    assert at.caption[0].value.startswith("North Division")
    for lbl, val in (("Claimant name", "Test Person"), ("Employee ID", "E1/2"), ("Hospital / Diagnostic Centre", "Some Lab")):
        _text(at, lbl).set_value(val)
    at.run()
    [b for b in at.button if "Generate documents" in b.label][0].click().run()
    assert not at.exception
    docs = at.session_state["generated_docs"]
    memo, ns = _doc_text(docs["memo"]), _doc_text(docs["ns"])
    assert "North Division" in memo and "Dak Bhawan, Delhi - 110001" in memo
    assert "A-1/NORTH/MR/" in memo and "2- The Director, Postal Accounts" in memo and "3- O/c" in memo
    assert "New Delhi Central Division" not in memo               # nothing of the old office is left behind
    assert "Sr. Accountant" in ns and "Included in LB012" in ns
    # both downloads stay on screen after generating
    labels = [b.label for b in at.get("download_button")]
    assert any("Sanction Memo" in x for x in labels) and any("NS/Checklist" in x for x in labels)


def test_default_office_reproduces_original_wording():
    at = _app(BILL)
    for lbl, val in (("Claimant name", "Test Person"), ("Employee ID", "E1")):
        _text(at, lbl).set_value(val)
    at.run()
    [b for b in at.button if "Generate documents" in b.label][0].click().run()
    memo = _doc_text(at.session_state["generated_docs"]["memo"])
    assert "O/o The Senior Superintendent of Post Offices," in memo
    assert "New Delhi Central Division" in memo and "Meghdoot Bhawan, New Delhi -110001" in memo
    assert "Director, NDHO for information and making payment to the claimant." in memo


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print("ok  ", name)
            except Exception as e:  # noqa: BLE001
                fails += 1
                print("FAIL", name, "-", repr(e))
    sys.exit(1 if fails else 0)
